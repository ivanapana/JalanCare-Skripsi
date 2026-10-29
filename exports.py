"""Agency forms and survey workbook. Field results are never inferred from photos."""
import io
from collections import defaultdict
from xml.sax.saxutils import escape
from dinas import GROUPS,FIELDS,META,NUMBERS,category_display,sta_m,hole_class

WAIT='Belum disurvei'
def text(value):return WAIT if value is None or value=='' else str(value)
def safe(value):return "'"+value if isinstance(value,str) and value.startswith(('=','+','-','@')) else value

def route_summary(rows):
    groups=defaultdict(list)
    for r in rows:
        d=r.get('surveiDinas') or {}
        if d.get('nomorRuas'):
            year=(d.get('tanggalSurvei') or '')[:4]
            groups[(d['nomorRuas'],year)].append(r)
    output=[]
    for (number,year),items in sorted(groups.items()):
        first=items[0].get('surveiDinas') or {};intervals=[];issue='';length=first.get('panjangRuasKm');names=set();widths=set();statuses=set()
        counts={k:0 for k in ('Baik','Sedang','Rusak Ringan','Rusak Berat')}
        for r in items:
            d=r.get('surveiDinas') or {};names.add(d.get('namaRuas'));widths.add(d.get('lebarRuasM'));statuses.add(d.get('statusFungsi'))
            category=category_display(r.get('kategoriKerusakan'));n=r.get('nilaiSDI')
            expected='Baik' if n is not None and n<=50 else 'Sedang' if n is not None and n<=100 else 'Rusak Ringan' if n is not None and n<=150 else 'Rusak Berat'
            if not year or not d.get('surveyor') or not d.get('namaRuas') or any(d.get(k) is None for k,*_ in FIELDS) or n is None or category!=expected:issue='Data survei/SDI belum lengkap atau belum konsisten'
            if d.get('panjangRuasKm')!=length:issue='Panjang ruas tidak konsisten'
            try:
                start,end=sta_m(d.get('staAwal')),sta_m(d.get('staAkhir'))
                if end-start!=200:raise ValueError()
                intervals.append((start,end))
                if category in counts:counts[category]+=end-start
            except (ValueError,TypeError):issue='STA segmen belum lengkap'
        if len(names)>1 or len(widths)>1 or len(statuses)>1:issue='Identitas ruas tidak konsisten'
        intervals.sort();cursor=0
        for start,end in intervals:
            if start!=cursor:issue='Cakupan STA terputus atau tumpang tindih; persentase tidak dihitung'
            cursor=end
        if length is None or length<=0 or abs(cursor-length*1000)>0.001:issue=issue or 'Cakupan belum seluruh ruas; persentase tidak dihitung'
        percentages=[None]*4 if issue else [round(counts[k]/(length*1000)*100,3) for k in counts]
        output.append([number,first.get('namaRuas'),year,first.get('lebarRuasM'),length,*percentages,first.get('aksesKeJalan'),issue or 'Seluruh ruas tercakup - hasil survei petugas'])
    return output

def workbook(rows,stream):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
    wb=Workbook();ws=wb.active;ws.title='Laporan Warga'
    keys=['idLaporan','timestamp','namaJalan','kabupaten','status','latitude','longitude','jumlahDeteksi','jumlahLubang','panjangCm','lebarCm','kedalamanCm','nilaiSDI','kategoriKerusakan','prioritasPenanganan','narasiLaporan','catatanPetugas']
    ws.append(['ID','Tanggal UTC','Nama Jalan/Patokan','Kabupaten/Kota','Status','Latitude','Longitude','Kotak Deteksi per Foto (bukan jumlah lubang segmen)','Jumlah Lubang Unik Foto (Petugas)','Panjang Lubang cm','Lebar Lubang cm','Kedalaman Lubang cm','SDI Surveyor','Kondisi Ruas (Petugas)','Prioritas (Petugas)','Narasi Awal','Catatan Petugas'])
    for r in rows:ws.append([safe(category_display(r.get(k)) if k=='kategoriKerusakan' else r.get(k)) for k in keys])
    analysis=wb.create_sheet('Analisis SDI')
    headers=['No.','Nomor Ruas','Nama Ruas','Panjang Jalan (km)','STA Awal','STA Akhir']+[label for _,label,_,_ in FIELDS]+['SDI Retak Luasan','SDI Retak Lebar','SDI Jumlah Lubang','SDI Bekas Roda','Nilai SDI (Petugas)','Kondisi','Kode','Tipe Perkerasan','ID Laporan','Tanggal Survei','Surveyor']
    analysis.append(headers)
    for index,r in enumerate(rows,1):
        d=r.get('surveiDinas') or {};category=category_display(r.get('kategoriKerusakan'))
        analysis.append([safe(v) for v in [index,d.get('nomorRuas'),d.get('namaRuas'),d.get('panjangRuasKm'),d.get('staAwal'),d.get('staAkhir'),*[d.get(k) for k,*_ in FIELDS],*[d.get(k) for k in ('sdiRetakLuas','sdiRetakLebar','sdiLubang','sdiBekasRoda')],r.get('nilaiSDI'),category,d.get('kodeKondisi'),'ASPAL' if d else None,r.get('idLaporan'),d.get('tanggalSurvei'),d.get('surveyor')]])
    rekap=wb.create_sheet('Rekap Ruas')
    rekap.append(['No. Ruas','Nama Ruas','Tahun Survei','Lebar Ruas (m)','Panjang Seluruh Ruas (km)','Baik (%)','Sedang (%)','Rusak Ringan (%)','Rusak Berat (%)','Akses ke Jalan N/P/K','Keterangan Cakupan'])
    for row in route_summary(rows):rekap.append([safe(v) for v in row])
    measurements=wb.create_sheet('Pengukuran Lubang')
    measurements.append(['ID Laporan','Diameter Ukur (cm)','Kedalaman Ukur (cm)','Klasifikasi Lubang Ukur','Jumlah Lubang Segmen 200 m (Survei)','Bukan Kotak YOLO'])
    for r in rows:
        d=r.get('surveiDinas') or {};measurements.append([safe(v) for v in [r.get('idLaporan'),d.get('diameterCm'),r.get('kedalamanCm'),hole_class(d.get('diameterCm'),r.get('kedalamanCm')),d.get('jumlahLubangSegmen'),'Hitungan foto tidak digunakan sebagai hitungan segmen']])
    legend=wb.create_sheet('Panduan Kode')
    legend.append(['Kelompok','Kolom','Kode','Arti'])
    for title,_,fields in GROUPS:
        for key,label,_,options in fields:
            for n,option in enumerate(options,1):legend.append([title,label,n,option])
    notes=wb.create_sheet('Catatan')
    for line in ['DRAF DIGITAL - mengacu kolom formulir PUPRD Sulut yang diberikan, bukan pengesahan dinas.','Kolom kosong berarti BELUM DISURVEI; kode 1 Tidak ada hanya berdasarkan pemeriksaan.','Ruas sampel 059 / Ir. Soekarno dan persentase sampel tidak disalin ke laporan.','SDI dan tahap analisis diisi surveyor; formula lengkap tidak terdapat dalam PDF sampel.','Kondisi: Baik <=50; Sedang >50 sampai 100; Rusak Ringan >100 sampai 150; Rusak Berat >150.','Rusak Ringan adalah tampilan dinas untuk kode Rusak yang sudah ada di aplikasi/skripsi.','Diameter <50 cm kecil; >=50 cm besar. Kedalaman <5 cm dangkal; >=5 cm dalam.','Ukuran satu lubang tidak menentukan kondisi seluruh ruas. Kode ukuran segmen dipilih surveyor.','Rekap persentase hanya dihitung jika seluruh STA 0 sampai panjang ruas lengkap, tanpa overlap, data survei lengkap dan satu tahun survei.','Ekspor mengikuti filter dashboard. Bila hanya sebagian ruas disaring, persentase seluruh ruas dikosongkan.','Kode kondisi dimasukkan surveyor sesuai standar dinas; aplikasi tidak menebak singkatannya.','Panduan gambar memakai jumlah /km; aplikasi mengikuti formulir segmen 200 m: 0, 1-2, 3-9, >=10.']:
        notes.append([line])
    for sheet in wb:
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='173F39');cell.alignment=Alignment(wrap_text=True,vertical='center')
        sheet.row_dimensions[1].height=46
        for column in sheet.columns:sheet.column_dimensions[column[0].column_letter].width=22
        for row in sheet.iter_rows(min_row=2):
            for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
        sheet.sheet_properties.pageSetUpPr.fitToPage=True;sheet.page_setup.orientation='landscape';sheet.page_setup.paperSize=sheet.PAPERSIZE_A3;sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0;sheet.print_title_rows='1:1'
    notes.column_dimensions['A'].width=130
    ws.column_dimensions['P'].width=75;ws.column_dimensions['Q'].width=45
    wb.save(stream)

def pdf_report(rows,stream):
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,PageBreak,KeepTogether
    from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from pathlib import Path
    import reportlab
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    fonts=Path(reportlab.__file__).parent/'fonts'
    if 'DinasRegular' not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont('DinasRegular',str(fonts/'Vera.ttf')))
        pdfmetrics.registerFont(TTFont('DinasBold',str(fonts/'VeraBd.ttf')))
    styles=getSampleStyleSheet()
    for style in styles.byName.values():style.fontName='DinasBold' if style.name in ('Title','Heading1','Heading2','Heading3') else 'DinasRegular'
    styles['Title'].fontSize=16;styles['Title'].leading=20
    styles['Heading3'].fontSize=10;styles['Heading3'].leading=13
    small=ParagraphStyle('FormSmall',fontName='DinasRegular',fontSize=6.5,leading=8)
    head=ParagraphStyle('FormHead',parent=small,fontName='DinasBold',alignment=1,fontSize=7,leading=9)
    def clean(v):return escape(text(v).replace('≥','>=').replace('–','-').replace('—','-'))
    def para(v,style=small):return Paragraph(clean(v),style)
    table_style=TableStyle([('BOX',(0,0),(-1,-1),.5,colors.black),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),3),('RIGHTPADDING',(0,0),(-1,-1),3),('TOPPADDING',(0,0),(-1,-1),3),('BOTTOMPADDING',(0,0),(-1,-1),3)])
    def field_box(label,options,code,width=124,right_code=None,paired=False):
        if paired:
            data=[[para('KR',head),para(label,head),para('KN',head)]]
            for n,opt in enumerate(options,1):data.append([para('[X]' if n==code else '[ ]'),para(str(n)+'. '+opt),para('[X]' if n==right_code else '[ ]')])
            box=Table(data,colWidths=[18,width-36,18])
        else:
            data=[[para(label,head)]]
            for n,opt in enumerate(options,1):data.append([para(('[X]' if n==code else '[ ]')+' '+str(n)+'. '+opt)])
            box=Table(data,colWidths=[width])
        box.setStyle(table_style);box.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e9eeee'))]));return box
    story=[]
    if not rows:story=[Paragraph('Tidak ada laporan sesuai filter.',styles['Title'])]
    for index,r in enumerate(rows):
        if index:story.append(PageBreak())
        d=r.get('surveiDinas') or {}
        story += [Paragraph('FORMULIR SURVEI KONDISI JALAN ASPAL',styles['Title']),Paragraph('DINAS PUPRD PROVINSI SULAWESI UTARA - DRAF DIGITAL',styles['Heading3']),para('Mengacu formulir yang diberikan. Belum merupakan pengesahan dinas. Kolom kosong: belum disurvei.'),Spacer(1,8)]
        meta=[['Nomor provinsi: 50','Nama provinsi: SULAWESI UTARA'],['Nomor ruas: '+text(d.get('nomorRuas')),'Nama ruas: '+text(d.get('namaRuas'))],['Dari patok Km: '+text(d.get('staAwal')),'Ke patok Km: '+text(d.get('staAkhir'))],['Status/fungsi: '+text(d.get('statusFungsi')),'Tanggal survei: '+text(d.get('tanggalSurvei'))],['Surveyor: '+text(d.get('surveyor')),'ID laporan: '+text(r.get('idLaporan'))]]
        top=Table([[para(v) for v in row] for row in meta],colWidths=[263,263]);top.setStyle(table_style);story += [top,Spacer(1,10)]
        columns=[]
        for title,_,fields in GROUPS:
            parts=[para(title,head),Spacer(1,5)]
            done=set()
            for key,label,_,options in fields:
                if key in done:continue
                if key.endswith('Kiri'):
                    other=key[:-4]+'Kanan';done.add(other)
                    parts += [field_box(label.rsplit(' ',1)[0],options,d.get(key),right_code=d.get(other),paired=True),Spacer(1,5)]
                else:parts += [field_box(label,options,d.get(key)),Spacer(1,5)]
            columns.append(parts)
        grid=Table([columns],colWidths=[131.5]*4);grid.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),0),('RIGHTPADDING',(0,0),(-1,-1),6),('BOX',(0,0),(-1,-1),.5,colors.black)]));story.append(grid)
        story += [Spacer(1,8),para('Ukuran: kecil diameter <50 cm; besar >=50 cm; dangkal kedalaman <5 cm; dalam >=5 cm.'),para('Status ruas: N Nasional; P Provinsi; M Kotamadya; K Kabupaten. KR kiri; KN kanan.'),para('Kolom tanpa [X] belum disurvei. Pemeriksaan retak, alur, bahu dan saluran dilakukan surveyor.'),PageBreak()]
        story += [Paragraph('LAMPIRAN HASIL PEMERIKSAAN DAN LAPORAN WARGA',styles['Heading1'])]
        gps=str(r['latitude'])+', '+str(r['longitude']) if r.get('latitude') is not None and r.get('longitude') is not None else None
        evidence=[('ID laporan',r.get('idLaporan')),('Lokasi foto',r.get('alamatLokasi') or r.get('namaJalan')),('GPS EXIF',gps),('Nilai SDI surveyor',r.get('nilaiSDI')),('Kondisi ruas',category_display(r.get('kategoriKerusakan'))),('Prioritas petugas',r.get('prioritasPenanganan')),('Diameter ukur cm',d.get('diameterCm')),('Kedalaman ukur cm',r.get('kedalamanCm')),('Klasifikasi lubang ukur',hole_class(d.get('diameterCm'),r.get('kedalamanCm'))),('Jumlah lubang segmen 200 m',d.get('jumlahLubangSegmen'))]
        evidence += [(label,d.get(k)) for k,label,_ in NUMBERS if k.startswith('sdi')]
        details=Table([[para(label),para(value)] for label,value in evidence],colWidths=[190,336]);details.setStyle(table_style);story += [details,Spacer(1,10),para('Ukuran satu lubang berbeda dari kondisi seluruh ruas. Tahap SDI adalah input surveyor, bukan hasil confidence YOLO.'),Spacer(1,8)]
        for photo in r.get('foto',[]):story.append(para(str(photo.get('angle'))+': '+str(len(photo.get('boxes',[])))+' kotak deteksi. Bukan jumlah lubang segmen.'))
        for title,key in [('Narasi awal saat pelaporan','narasiLaporan'),('Catatan petugas terbaru','catatanPetugas')]:
            story += [Spacer(1,8),Paragraph(title,styles['Heading3']),Paragraph(clean(r.get(key)).replace('\n','<br/>'),styles['Normal'])]
    def footer(canvas,doc):
        canvas.saveState();canvas.setFont('DinasRegular',7);canvas.drawString(34,22,'JalanCare | Draf digital - verifikasi surveyor diperlukan');canvas.drawRightString(A4[0]-34,22,str(doc.page));canvas.restoreState()
    SimpleDocTemplate(stream,pagesize=A4,rightMargin=34,leftMargin=34,topMargin=28,bottomMargin=38).build(story,onFirstPage=footer,onLaterPages=footer)

def compact_workbook(rows,stream):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    wb=Workbook();ws=wb.active;ws.title='Laporan Lubang'
    ws.append(['ID','Tanggal UTC','Jalan/Patokan','Lokasi','Status','Latitude EXIF','Longitude EXIF','Kotak Deteksi (bukan jumlah lubang unik)','Narasi Awal','Catatan Petugas'])
    for r in rows:ws.append([safe(r.get(k)) for k in ('idLaporan','timestamp','namaJalan','alamatLokasi','status','latitude','longitude','jumlahDeteksi','narasiLaporan','catatanPetugas')])
    survey=wb.create_sheet('Verifikasi Petugas')
    survey.append(['ID','Jumlah Lubang Unik (Petugas)','Panjang cm','Lebar cm','Kedalaman cm','Nilai SDI (Petugas)','Kategori (Petugas)','Prioritas (Petugas)'])
    for r in rows:survey.append([safe(category_display(r.get(k)) if k=='kategoriKerusakan' else r.get(k)) for k in ('idLaporan','jumlahLubang','panjangCm','lebarCm','kedalamanCm','nilaiSDI','kategoriKerusakan','prioritasPenanganan')])
    notes=wb.create_sheet('Catatan')
    for line in ['Draf awal otomatis berisi indikasi lubang dan lokasi, bukan survei seluruh ruas.','Kategori kerusakan ditetapkan surveyor sesuai Bab1.5.2 poin14 skripsi Ivana.','Pengukuran dan SDI diisi petugas apabila hasil survei tersedia. Kolom kosong berarti belum dinilai.','Confidence dan jumlah kotak foto tidak menentukan kategori, dimensi fisik atau jumlah lubang unik.','Tidak memuat isian survei retak, bahu, saluran dan perkerasan seluruh ruas.','Rusak Ringan adalah label dinas untuk kode Rusak yang tetap disimpan agar kompatibel dengan data lama.']:
        notes.append([line])
    for sheet in wb:
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='173F39');cell.alignment=Alignment(wrap_text=True)
        sheet.row_dimensions[1].height=42
        for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=25
        for row in sheet.iter_rows(min_row=2):
            for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
    ws.column_dimensions['I'].width=75;notes.column_dimensions['A'].width=120;wb.save(stream)

def compact_pdf(rows,stream):
    from pathlib import Path
    import reportlab
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,PageBreak
    from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    fonts=Path(reportlab.__file__).parent/'fonts'
    if 'DinasRegular' not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont('DinasRegular',str(fonts/'Vera.ttf')))
        pdfmetrics.registerFont(TTFont('DinasBold',str(fonts/'VeraBd.ttf')))
    styles=getSampleStyleSheet();normal=ParagraphStyle('CompactNormal',fontName='DinasRegular',fontSize=9,leading=13)
    heading=ParagraphStyle('CompactHeading',fontName='DinasBold',fontSize=13,leading=18,spaceAfter=9)
    def para(value):return Paragraph(escape('Belum tersedia' if value is None or value=='' else str(value)).replace('–','-').replace('—','-').replace('\n','<br/>'),normal)
    story=[]
    if not rows:story.append(para('Tidak ada laporan sesuai filter.'))
    for index,r in enumerate(rows):
        if index:story.append(PageBreak())
        story += [Paragraph('LAPORAN INDIKASI LUBANG JALAN',heading),para('Draf awal pelaporan. Kategori kondisi jalan ditetapkan surveyor Dinas PUPRD.'),Spacer(1,10)]
        pairs=[('ID','idLaporan'),('Tanggal UTC','timestamp'),('Jalan/Patokan','namaJalan'),('Lokasi','alamatLokasi'),('Status laporan','status'),('Latitude EXIF','latitude'),('Longitude EXIF','longitude'),('Kotak deteksi (bukan lubang unik)','jumlahDeteksi'),('Kategori hasil pemeriksaan petugas','kategoriKerusakan'),('Catatan petugas','catatanPetugas')]
        data=[[para(label),para(category_display(r.get(key)) if key=='kategoriKerusakan' else r.get(key))] for label,key in pairs]
        table=Table(data,colWidths=[200,315]);table.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.4,colors.HexColor('#ccd8d4')),('VALIGN',(0,0),(-1,-1),'TOP'),('BACKGROUND',(0,0),(0,-1),colors.HexColor('#edf4f1')),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]));story += [table,Spacer(1,14)]
        fields=[(label,r.get(k)) for k,label in [('jumlahLubang','Jumlah lubang unik'),('panjangCm','Panjang cm'),('lebarCm','Lebar cm'),('kedalamanCm','Kedalaman cm'),('nilaiSDI','Nilai SDI surveyor'),('prioritasPenanganan','Prioritas petugas')] if r.get(k) is not None and r.get(k)!='']
        if fields:
            story += [Paragraph('Hasil pemeriksaan yang tersedia',heading)]
            for label,value in fields:story.append(para(label+': '+str(value)))
            story.append(Spacer(1,12))
        story += [Paragraph('Narasi awal tersimpan',heading),para(r.get('narasiLaporan'))]
    def footer(canvas,doc):
        canvas.saveState();canvas.setFont('DinasRegular',7);canvas.drawString(40,22,'JalanCare | Laporan lubang - verifikasi petugas');canvas.drawRightString(A4[0]-40,22,str(doc.page));canvas.restoreState()
    SimpleDocTemplate(stream,pagesize=A4,rightMargin=40,leftMargin=40,topMargin=35,bottomMargin=40).build(story,onFirstPage=footer,onLaterPages=footer)

def build_export(rows,fmt):
    stream=io.BytesIO()
    if fmt in ('xlsx','xlsx-agency'):
        (workbook if fmt=='xlsx-agency' else compact_workbook)(rows,stream);mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    else:
        (pdf_report if fmt=='pdf-agency' else compact_pdf)(rows,stream);mime='application/pdf'
    stream.seek(0);return stream,mime
