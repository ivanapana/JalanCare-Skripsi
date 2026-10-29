"""Run without Flask/YOLO: python -m unittest discover -s tests -p test_dinas_unit.py."""
import io,sys,unittest
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from dinas import FIELDS,parse_survey,hole_class,category_display,sta_m
from exports import build_export,route_summary
from narrative import format_report
from i18n import translate

class DinasTests(unittest.TestCase):
    def test_unknown_is_not_absent(self):
        d=parse_survey({'dinas_form':'1'},None)
        self.assertTrue(all(d[k] is None for k,*_ in FIELDS))
        self.assertIsNone(d['jumlahLubangSegmen'])
        self.assertEqual(parse_survey({},d),d)
    def test_segment_counts_use_field_survey(self):
        base={'dinas_form':'1','dinas_staAwal':'0+000','dinas_staAkhir':'0+200'}
        for count,code in [(0,1),(1,2),(2,2),(3,3),(9,3),(10,4)]:
            self.assertEqual(parse_survey({**base,'dinas_jumlahLubangSegmen':str(count)},None)['jumlahKode'],code)
        with self.assertRaises(ValueError):parse_survey({**base,'dinas_jumlahLubangSegmen':'3','dinas_jumlahKode':'2'},None)
        with self.assertRaises(ValueError):parse_survey({**base,'dinas_staAkhir':'0+400'},None)
        with self.assertRaises(ValueError):parse_survey({'dinas_form':'1','dinas_jumlahLubangSegmen':'3'},None)
    def test_size_boundaries_are_physical(self):
        for diameter,depth,expected in [(49.9,4.9,'Kecil-dangkal'),(49.9,5,'Kecil-dalam'),(50,4.9,'Besar-dangkal'),(50,5,'Besar-dalam')]:self.assertEqual(hole_class(diameter,depth),expected)
        self.assertIsNone(hole_class(None,5));self.assertEqual(category_display('Rusak'),'Rusak Ringan')
    def test_invalid_codes_and_values(self):
        for values in [{'dinas_susunan':'3'},{'dinas_tanggalSurvei':'2026-02-30'},{'dinas_diameterCm':'nan'},{'dinas_panjangRuasKm':'0'},{'dinas_statusFungsi':'X'},{'dinas_staAwal':'bad'},{'dinas_jumlahLubangSegmen':'2.5'}]:
            with self.assertRaises(ValueError):parse_survey({'dinas_form':'1',**values},None)
        self.assertEqual(sta_m('1+200'),1200)
        self.assertEqual(translate('SDI tahap bekas roda dan nilai SDI akhir harus sama.','en'),'The SDI after rutting and final SDI value must match.')
    def row(self,start='0+000',end='0+200',sdi=120,category='Rusak'):
        d={'nomorRuas':'UJI','namaRuas':'Ruas Uji','tanggalSurvei':'2026-10-14','surveyor':'Surveyor Uji','panjangRuasKm':.4,'lebarRuasM':6,'statusFungsi':'P','staAwal':start,'staAkhir':end,**{k:1 for k,*_ in FIELDS}}
        return {'idLaporan':'uji','nilaiSDI':sdi,'kategoriKerusakan':category,'surveiDinas':d,'namaJalan':'Titik foto','foto':[],'latitude':1,'longitude':124,'timestamp':'2026-10-14','kabupaten':'Minahasa','narasiLaporan':'Uji','keterangan':''}
    def test_percent_requires_full_nonoverlapping_survey(self):
        one=self.row();two=self.row('0+200','0+400',30,'Baik')
        self.assertEqual(route_summary([one])[0][5:9],[None]*4)
        self.assertEqual(route_summary([one,one])[0][5:9],[None]*4)
        self.assertEqual(route_summary([one,two])[0][5:9],[50,0,50,0])
        two['surveiDinas']['retakLebar']=None
        self.assertEqual(route_summary([one,two])[0][5:9],[None]*4)
    def test_excel_fields_blank_and_formula_safe(self):
        from openpyxl import load_workbook
        r=self.row();r['surveiDinas']['surveyor']='=HYPERLINK("bad")';r['surveiDinas']['susunan']=None
        stream,mime=build_export([r],'xlsx-agency');wb=load_workbook(stream)
        self.assertEqual(wb.sheetnames,['Laporan Warga','Analisis SDI','Rekap Ruas','Pengukuran Lubang','Panduan Kode','Catatan'])
        self.assertIsNone(wb['Analisis SDI']['G2'].value)
        self.assertEqual(wb['Laporan Warga']['N2'].value,'Rusak Ringan')
        self.assertTrue(wb['Analisis SDI'].cell(2,39).value.startswith("'="))
        self.assertIsNone(wb['Rekap Ruas']['F2'].value)
    def test_narrative_fields_do_not_invent_survey(self):
        r=self.row();r.pop('surveiDinas');r['nilaiSDI']=None;r['kategoriKerusakan']=None;r['foto']=[{'angle':'foto1','boxes':[{'confidence':.73,'xyxy':[50,50,200,200]}]}]
        s=format_report(r,'Uji','Periksa lapangan')
        self.assertIn('DRAF LAPORAN INDIKASI LUBANG JALAN',s)
        self.assertNotIn('Retak-retak',s)
        self.assertNotIn('Ir. Soekarno',s);self.assertLess(len(s),6500)
    def test_pdf_contains_agency_fields(self):
        from pypdf import PdfReader
        stream,_=build_export([self.row()],'pdf-agency');reader=PdfReader(stream)
        first=reader.pages[0].extract_text()
        for label in ['FORMULIR SURVEI KONDISI JALAN ASPAL','Nomor ruas','Retak-retak','Kondisi bahu','Kerusakan tepi']:
            self.assertIn(label,first)
        self.assertIn('Rusak Ringan',' '.join(p.extract_text() for p in reader.pages))

class SimplifiedScopeTests(unittest.TestCase):
    def test_main_ui_only_pothole_verification(self):
        html=Path(__file__).resolve().parents[1].joinpath('templates/detail.html').read_text()
        self.assertNotIn("include 'survey_dinas",html)
        self.assertIn('Pengukuran lubang oleh surveyor (opsional)',html)
        self.assertIn('Kategori hasil pemeriksaan petugas',html)
        self.assertNotIn('dinas_form',html)
    def test_simple_export_has_no_full_road_survey(self):
        from openpyxl import load_workbook
        r=DinasTests().row()
        stream,_=build_export([r],'xlsx');wb=load_workbook(stream)
        self.assertEqual(wb.sheetnames,['Laporan Lubang','Verifikasi Petugas','Catatan'])
        self.assertEqual(wb['Verifikasi Petugas']['G2'].value,'Rusak Ringan')
        self.assertEqual(r['kategoriKerusakan'],'Rusak')
    def test_pdf_is_a_pothole_report(self):
        from pypdf import PdfReader
        r=DinasTests().row();stream,_=build_export([r],'pdf');pdf=PdfReader(stream)
        words=' '.join(p.extract_text() for p in pdf.pages)
        self.assertIn('LAPORAN INDIKASI LUBANG JALAN',words)
        self.assertNotIn('Permukaan Perkerasan',words)
        self.assertIn('Rusak Ringan',words)

if __name__=='__main__':unittest.main()
